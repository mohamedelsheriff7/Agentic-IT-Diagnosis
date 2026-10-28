"""Build the human patch history from the retained verification records."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'reports'
OUT.mkdir(parents=True, exist_ok=True)
pages = []


def page(*blocks):
    pages.append(blocks)


def h(text): return ('h', text)
def sub(text): return ('sub', text)
def p(text): return ('p', text)
def b(text): return ('b', text)
def table(headers, rows, widths): return ('table', headers, rows, widths)


page(
    ('title', 'ISYS Diagnostic Agent Patch History'),
    p('Development history through version 6'),
    p('25 October 2026  |  Workstream D  |  Phase 5 Diagnostic Agent and Orchestration'),
    h('Overview'),
    p('The work progressed from a locally hosted Python diagnostic service to a native LangGraph agent for watsonx Orchestrate. Early patches repaired packaging, credentials and diagnostic-state consistency. Later patches made replies conversational, interpreted reported observations, and aligned the diagnostic loop with the Phase 5 requirements.'),
    p('Version 6 is the current consolidated patch. It passed 151 automated tests, 14 live backend checks, isolated package loading and a localhost Wi-Fi conversation. Its WXO import has not yet been confirmed by the user. The original version 5 ZIP remains unchanged for rollback.'),
    table(['Stage', 'Tests recorded', 'Principal change'], [
        ['Local API', '46', 'Authenticated localhost service and backend integration'],
        ['Version 1', '52', 'First native LangGraph package and message-based state'],
        ['Version 2', '52 / 54 / 55', 'Self-contained import and WXO credential binding'],
        ['Version 3', '59', 'Repair obsolete outage escalation state'],
        ['Version 4', '64', 'Reopen excluded hypotheses when evidence conflicts'],
        ['Version 4.1', '66', 'Clear stale questions and improve handoff traces'],
        ['Version 4.2', '90', 'Canonical taxonomy and post-handoff acknowledgements'],
        ['Version 5', '103', 'Natural replies with persistent conversation checkpoints'],
        ['Version 5.1', '140', 'Interpret DNS Wi-Fi and printer observations'],
        ['Version 6', '151', 'Consolidated fixes against the seven version 5 tests'],
    ], [0.9, 1.1, 5.0]),
    p('Counts are suite totals recorded at each stage, not additional tests to add together. A passed integration check establishes the tested behavior; it does not measure diagnostic accuracy across the full project dataset.'),
)

page(
    h('Local service and native import'),
    sub('Foundation work before numbered native releases'),
    p('The premade Python agent was exposed through a FastAPI service, with localhost startup scripts, authentication and an OpenAPI 3.0.3 export. Ticket investigation, answer submission, trace retrieval and explicit resolution confirmation were exposed as separate operations. Session data was persisted per ticket with paths independent of the current working directory.'),
    p('The backend adapter used the Workstream C URL and X-API-Key authentication. The reviewed Workstream B knowledge snapshot was copied into the project without changing its approval flags. Backend user responses were recorded with the ticket affected-person ID and the actual evidence IDs returned by the backend.'),
    p('Early safeguards covered secret redaction while preserving identifiers, explicit credential checks, unique action IDs, correct source and enum values, no fabricated tool-check evidence, exact resolution step IDs and success confirmation, and no resolution claim after a rejected request. An explicit empty backend key did not fall back silently to an environment value. Access to another application was a clue rather than conclusive exclusion of all account or password problems. Arabic retrieval and persisted traces were also tested. The local/API suite reached 46 passing tests.'),
    sub('Version 1 native LangGraph conversion'),
    p('The user requested native agent import rather than registering the agent as an OpenAPI tool. The first native package supplied agent.yaml, pinned dependencies and an uncompiled create_agent factory. Workflow nodes handled investigation, answers, explicit confirmation, help and responses. A local /langgraph/chat route used the same factory with SQLite checkpoints.'),
    p('Diagnostic snapshots travelled in structured assistant messages so a fresh factory could resume from messages alone. The package included synthetic knowledge and historical context, excluded credentials and evaluation answers, and passed the original extraction and streaming checks. The native suite reached 52 tests. Actual WXO execution then exposed a relocated-entrypoint import failure: No module named diagnostic_agent.'),
    sub('Version 2 packaging and credential rebuilds'),
    p('The entrypoint embedded an allowlisted archive of core modules and knowledge, extracted it into a temporary runtime directory, and adjusted the import path. Internal engine and history imports were made package-relative. An isolated subprocess received only agent.py in an unrelated working directory; factory construction, multi-turn resume and asynchronous streaming passed.'),
    p('Successive v2 rebuilds then repaired credential handling. The first missing-key message pointed only to a local environment variable. Credential resolution was extended to accept the named WXO connection at factory creation or invocation and to support key rotation without accepting unrelated credentials. Missing-configuration replies identified the required team connection.'),
    p('The user confirmed the real connection app ID was isys_it_backend. The primary credential became isys_it_backend_api_key, with workstream_c_api_key retained as a legacy alias. The recorded suite progressed from 52 to 54 and then 55 tests. These rebuilds reused the v2 filename; the retained archive is the final v2 build, not a separate archive for each intermediate credential fix.'),
)

page(
    h('Outage state and evidence reconciliation'),
    sub('Version 3 obsolete escalation state'),
    p('WXO reached the backend but retained an outage escalation after the corresponding VPN_OUTAGE hypothesis had been eliminated. Version 3 cleared only that superseded outage decision, continued individual investigation and repaired previously saved stale snapshots. Mandatory human restrictions remained in force.'),
    p('History agreement became NOT_AVAILABLE when no diagnosis existed to compare. Structural trace enums were preserved during redaction. The exact pasted snapshot was added as a regression fixture. The suite reached 59 tests, and a live VPN check returned ASK_MORE with the colleague-connectivity question.'),
    sub('Version 4 contradictory evidence'),
    p('The answer no supported a shared VPN problem, but the excluded VPN_OUTAGE hypothesis was not restored. Version 4 introduced archived candidates and reopened an excluded hypothesis as plausible when later observations supported it. Supporting and contradicting evidence IDs were retained separately, including when older snapshots lacked archived metadata.'),
    p('A later contradiction could invalidate an established conclusion. Active and eliminated candidate sets remained disjoint. Shared-service conflicts produced scoped screening and Supervisor review, without declaring an outage or claiming a global incident or link had been created. The suite reached 64 tests; isolated import and a live investigate-then-no sequence passed.'),
    sub('Version 4 point 1 handoff cleanup'),
    p('Handoffs still displayed an answered question and incomplete screening explanations. Version 4.1 returned next_question only for ASK_MORE, replaced stale questions with the real evidence-reconciliation gap, and added the completed recent-ticket screening result to the rationale. Zero matching recent tickets did not become proof of an outage.'),
    p('Identical consecutive outage trace events were suppressed. Public machine-readable evidence requirement markers were preserved by redaction while explicit secret assignments remained masked. The suite reached 66 tests.'),
    sub('Version 4 point 2 canonical values and acknowledgements'),
    p('A broader live matrix exposed unsupported UNKNOWN_CAUSE identifiers in unmatched-ticket paths. The patch used UNKNOWN_ISSUE and restricted backend candidate and elimination lists to the packaged canonical taxonomy. The initial live matrix passed 4 of 7 cases; the corrected run passed every case.'),
    p('Replies such as Thanks after a handoff became status requests rather than new diagnostic answers. They reused the recorded recommendation without duplicate evidence, worklogs or backend actions; a different ticket could not inherit another ticket handoff. The suite reached 90 tests.'),
    p('The user confirmed the complete two-turn VPN conflict and handoff sequence in WXO v4.2. The transcript retained conflicting evidence, cleared the answered question and did not claim global-ticket actions. This is tenant evidence for that scenario, not proof that every scenario ran in WXO.'),
)

page(
    h('Conversational replies and observation handling'),
    sub('Version 5 natural replies with retained state'),
    p('The agent had been exposing its full structured JSON state in chat. Version 5 rendered a short natural-language reply while retaining the diagnostic payload in AIMessage.additional_kwargs under isys_diagnostic. Legacy JSON replies remained readable, and structured output remained available as an explicit debugging option.'),
    p('To survive hosts that retain visible text but discard metadata, the agent also saved sanitized conversation checkpoints as existing backend NOTE worklogs. Checkpoints used unique references and action IDs, compressed chunks, checksums and bounded decoding. A short Ticket and Ref receipt in the visible reply identified the state to restore. These notes were not tool evidence.'),
    p('Checkpoint failures did not claim safe continuation. Acknowledgements added no duplicate state-writing actions, and resolved or closed tickets were not checkpointed. The localhost chat response kept structured fields for integrations and added a reply field for the visible text. The suite reached 103 tests.'),
    p('A live password conversation passed initial question, answer and acknowledgement using a fresh graph with only visible text on each turn. The user subsequently confirmed natural replies and the two-turn password scenario in WXO v5.'),
    sub('Version 5 point 1 observation interpretation'),
    p('Saved DNS, Wi-Fi and printer conversations showed that sentence answers were stored with empty support and contradiction mappings. The original interpreter largely recognized bare yes or no. After exhausting its small question bank, it returned the same generic human handoff.'),
    p('Version 5.1 added question-scoped interpretation for DNS, Wi-Fi, printer, synchronization and certificate observations. Distinct reported facts were retained with their original USER_RESPONSE evidence IDs. Combined answers could satisfy more than one question. Either-or answers received one clarification, and unknown or contradictory statements remained uncertain.'),
    p('A working second device no longer eliminated every Wi-Fi hypothesis. Legacy unmapped answers could be reinterpreted without creating new user evidence. A fresh preview could restore the checkpoint matching the backend latest diagnostic result rather than relying on worklog ordering or timestamps.'),
    p('Handoffs explained reported observations, a tentative hypothesis, the suggested support team and a specific missing check. The suite reached 140 tests; six live checks covered recovery and fresh text-only conversations for the three reported issues. The user later confirmed the improved DNS and Wi-Fi wording in WXO v5.1.'),
    p('Version 5.1 still required supporting tool evidence before confirming any diagnosis. The user then clarified that Phase 5 allowed sufficient reported evidence to support a diagnosis and that diagnosis, routing and resolution eligibility were separate decisions. That clarification drove the consolidated v6 change.'),
)

page(
    h('Version 5 batch findings and the consolidated response'),
    p('The user returned to the unchanged v5 package and supplied seven complete test transcripts before another patch was built. This created a common baseline for identifying related failures together.'),
    table(['Test', 'Observed version 5 behavior', 'Consolidated correction'], [
        ['1', 'Bare no plus account evidence reached PASSWORD_EXPIRED.', 'Keep this as the working comparison case.'],
        ['2', 'Equivalent sentence answer led to a generic handoff.', 'Map the account-access observation and match test 1.'],
        ['3', 'Asked again about external sites after a combined DNS answer.', 'Extract all supplied facts and skip answered questions.'],
        ['4', 'A working second device caused immediate handoff.', 'Repair the Wi-Fi exclusion and investigate the device/network pairing.'],
        ['5', 'Shared printer scope and failed reachability were not used clearly.', 'Interpret both observations and explain the supported problem pattern.'],
        ['6', 'Local configuration was not interpreted.', 'Retain the report and collect the unanswered VPN scope information.'],
        ['7', 'Unknown service/error immediately produced a generic handoff.', 'Ask which application or device was being used before abandoning scope collection.'],
    ], [0.45, 3.0, 3.55]),
    sub('Version 6 answer handling and question selection'),
    p('The patch combined the earlier observation fixes with account-access sentence interpretation and VPN configuration answers. Question-scoped facts preserved the original evidence IDs and source types. Commands, hedged statements, repeated assertions and unclear alternatives did not become fabricated checks or successful actions.'),
    p('Legacy state repair reversed the incorrect Wi-Fi exclusion and recovered meaningful unanswered questions. It also distinguished provisional questions selected internally from questions actually answered by the user. A healthy service observation alone did not replace the missing colleague-connectivity answer in the VPN case.'),
    sub('Version 6 diagnosis scope and confidence'),
    p('The blanket requirement for a diagnostic tool in every confirmed diagnosis was replaced with explicit reported-observation patterns for supported question families: DNS, Wi-Fi, printer reachability, VPN client issues, synchronization delay and certificate expiry. Matching knowledge and the required distinct observations were necessary; repetition of one statement was insufficient.'),
    p('These patterns established a symptom-level diagnosis with MEDIUM confidence and a routing recommendation. They did not identify an unverified faulty component or authorize a repair. Conflicting facts invalidated the reported pattern. Machine checks remained necessary where the particular evidence or resolution conditions required them.'),
)

page(
    h('Version 6 knowledge sources and verification'),
    sub('Diagnostic guidance and backend references'),
    p('Draft reviewed articles were permitted to guide questions and diagnosis in the synthetic diagnostic review mode, without changing their approval flags or activating draft routing or resolution policy. Guidance was limited to relevant current candidates. The approved-only behavior remained the default for callers outside that review mode.'),
    p('Live validation exposed a further integration error: bundled KB-026 was absent from the backend article registry, but the agent sent it as a backend knowledge reference. The patch separated bundled guidance from registered backend foreign keys. Bundled IDs, versions and approval status were retained in knowledge_sources and the diagnostic trace, while unverified IDs were excluded from backend reference arrays.'),
    p('The original diagnosis, evidence, confidence, recommended action and rationale remained available as structured data. The result also identified whether a diagnosis represented a reported symptom pattern, a corroborated diagnosis, or an undetermined case. User-facing messages continued to be conversational.'),
    sub('Saved transcripts and fresh live conversations'),
    table(['Test', 'Recovered original ticket', 'Fresh conversation'], [
        ['1', 'PASSWORD_EXPIRED and ROUTE', 'Same result with short answers'],
        ['2', 'PASSWORD_EXPIRED and ROUTE', 'Same result with sentence answers'],
        ['3', 'DNS_FAILURE and ROUTE', 'Combined answer used without repetition'],
        ['4', 'ASK_MORE about another network', 'WIFI_FAILURE and ROUTE after both facts'],
        ['5', 'PRINTER_OFFLINE and ROUTE', 'Shared scope and reachability used'],
        ['6', 'ASK_MORE about colleague connectivity', 'VPN_CLIENT_ISSUE and ROUTE after scope and error'],
        ['7', 'ASK_MORE about application or device', 'Continued focused scope collection'],
    ], [0.45, 3.0, 3.55]),
    p('All 14 live checks passed. Recovered conversations retained their original evidence IDs. Fresh turns discarded custom graph state and message metadata, exercising visible-text checkpoint recovery. No assignments or resolutions were performed. An intermittent transport failure interrupted the run; a healthy backend probe and a resumed run completed the pending checks.'),
    p('The final suite passed 151 tests. The v6 ZIP also passed isolated single-file loading, fresh-factory message resume, a natural printer conversation and asynchronous streaming. A localhost Wi-Fi smoke check continued past yes to the alternate-network question and then returned WIFI_FAILURE / ROUTE. The v5 artifact hash was checked and remained unchanged.'),
)

page(
    h('Current implementation and retained boundaries'),
    p('The current deliverable is a native LangGraph workflow hosted by WXO, with a packaged deterministic diagnostic engine and bundled knowledge. It is not an autonomous LLM tool-execution loop. WXO obtains backend credentials through the isys_it_backend team connection; local credentials are not uploaded in the ZIP.'),
    b('Conversation state includes the ticket, candidates, evidence IDs, answered questions, exclusions, source references, progress and the latest diagnostic result. Message metadata and backend checkpoint notes support recovery across turns and fresh previews.'),
    b('Diagnostic checks currently cover account status, application entitlement and service health. The patch did not implement DNS probes, device management, printer control or new production monitoring integrations. Backend check results come from synthetic project facts.'),
    b('Diagnosis and resolution are separate. A supported diagnosis may recommend routing while the underlying component fault remains unverified. Resolution still requires an eligible article, its exact conditions and steps, actual user confirmation, a matching diagnostic record and backend success. This batch did not exercise a successful end-to-end resolution walkthrough.'),
    b('The Diagnostic Agent recommends routing or escalation. Supervisor or human handling remains responsible for assignments and global-incident operations. A human-review request is not approval, and a recommendation is not a completed action.'),
    b('The current reported-observation patterns and question interpretation are bounded rules. The completed tests do not establish full diagnostic accuracy, broad free-language understanding, or production readiness across the complete evaluation dataset.'),
    h('Verification evidence and tenant status'),
    table(['Evidence', 'What it establishes'], [
        ['Automated regressions', 'Specified transitions, interpretation rules, evidence provenance and safeguards under the tested fixtures.'],
        ['Live synthetic backend checks', 'Authentication, schema/reference acceptance, persisted results and continuation for the exercised scenarios.'],
        ['Isolated package checks', 'The relocated agent.py can load its embedded code and run the tested native graph interfaces.'],
        ['User WXO transcripts', 'Actual tenant behavior for the supplied v4.2, v5 and v5.1 conversations; seven v5 cases formed the v6 baseline.'],
        ['Version 6 tenant execution', 'Package delivered and locally verified; actual v6 WXO import and user execution remain unconfirmed.'],
    ], [1.7, 5.3]),
    p('Earlier verification files contain stage-specific statements such as tenant execution not yet verified. Those statements describe that release checkpoint. Later user transcripts supersede them for the particular scenario demonstrated; they do not automatically validate every other scenario.'),
)

inventory = json.loads((ROOT / 'docs/patch_history_inventory.json').read_text(encoding='utf-8'))
page(
    h('Artifact inventory and source records'),
    p('Retained ZIPs are stored in dist. The following SHA256 prefixes identify the archives inspected for this report. Full hashes, resource hashes and changed-resource lists are recorded in docs/patch_history_inventory.json. Version 1 is the original unnumbered native ZIP.'),
    table(['Version', 'Archive filename', 'SHA256 prefix'], [
        [x['version'].replace('v',''), x['file'], x['sha256'][:12]] for x in inventory
    ], [0.9, 4.6, 1.5]),
    sub('Primary development records'),
    p('LOCAL_VERIFICATION.md records the localhost foundation, native import stages, credential rebuilds, outage-state patches and test counts. CONSOLIDATED_DEBUGGING_REPORT.md records v4.2 validation, taxonomy repairs and the seven-case live matrix. These are historical checkpoints.'),
    p('conversational_backend_verification.json records the v5 text-only password continuation. natural_language_backend_verification.json records the v5.1 DNS, Wi-Fi and printer checks. v5_batch_test_tickets.json identifies the seven untouched baseline tickets and the original v5 checksum. The user transcripts and tests/fixtures/v5_batch_reports.json preserve the actual v5 failures.'),
    p('V6_CONSOLIDATED_PATCH.md and v6_live_batch_results.json document the consolidated change and its 14 completed live checks. test_v5_batch_patch.py replays the seven saved cases. test_natural_observations.py and test_conversation.py cover observation parsing, checkpoint recovery and conversation behavior. verify_langgraph_package.py checks the built native artifact in isolation.'),
    sub('Principal implementation areas'),
    p('wxo_langgraph/agent.py owns the native factory, credential injection and conversational state transport. engine.py owns candidates, questions and diagnostic transitions. answers.py interprets reported facts and supported patterns. integrated.py records backend evidence and results. conversation.py renders replies and saves/restores checkpoints. knowledge.py handles retrieval. models.py and session_store.py retain the diagnostic state.'),
    p('No API keys, credential values, environment-file contents or private user data are included in this report. Test ticket identifiers refer to the synthetic project records.'),
)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.75)
sec.footer_distance = Inches(0.3)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.07
for name, size in [('Title', 25), ('Heading 1', 17), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.color.rgb = 'Arial', Pt(size), RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True
    if name != 'Title': style.font.bold = True
footer = sec.footer.paragraphs[0]
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
run = footer.add_run('ISYS Diagnostic Agent Patch History\t')
run.font.name, run.font.size, run.font.color.rgb = 'Arial', Pt(8), RGBColor(100,100,100)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
md = []
for index, blocks in enumerate(pages):
    if index: doc.add_page_break()
    for block in blocks:
        kind = block[0]
        if kind in {'title','h','sub'}:
            style = {'title':'Title','h':'Heading 1','sub':'Heading 2'}[kind]
            doc.add_paragraph(block[1], style)
            md.extend([{'title':'# ','h':'## ','sub':'### '}[kind] + block[1], ''])
        elif kind in {'p','b'}:
            para = doc.add_paragraph(block[1], 'List Bullet' if kind == 'b' else 'Normal')
            para.paragraph_format.widow_control = True
            md.extend([('- ' if kind == 'b' else '') + block[1], ''])
        elif kind == 'table':
            headers, rows, widths = block[1:]
            t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.autofit=False
            for i,w in enumerate(widths): t.columns[i].width=Inches(w)
            for i,text in enumerate(headers): t.rows[0].cells[i].text = text
            repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                cells=t.add_row().cells
                for i,text in enumerate(row): cells[i].text=str(text)
            for ri,row in enumerate(t.rows):
                trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
                for ci,cell in enumerate(row.cells):
                    cell.width=Inches(widths[ci])
                    if ri == 0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9EDF2'); cell._tc.get_or_add_tcPr().append(shade)
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(4); para.paragraph_format.space_before=Pt(4)
                        para.paragraph_format.line_spacing=1.0
                        for run in para.runs:
                            run.font.name='Arial'; run.font.size=Pt(9.5); run.bold=(ri==0)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(headers)+' |')
            md.append('| '+' | '.join('---' for _ in headers)+' |')
            md.extend('| '+' | '.join(str(v) for v in row)+' |' for row in rows)
            md.append('')
doc.core_properties.title = 'ISYS Diagnostic Agent Patch History'
doc.core_properties.subject = 'Patch history through version 6'
doc.core_properties.author = 'ISYS project development record'
# The bundled default template can carry a blue Title paragraph border.
# Typography and whitespace provide the title separation for this report.
for container in (doc.styles.element, doc.element.body):
    for border in list(container.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT / 'ISYS_Diagnostic_Agent_Patch_History.docx')
(OUT / 'ISYS_Diagnostic_Agent_Patch_History.md').write_text('\n'.join(md), encoding='utf-8')
print(json.dumps({'docx':str(OUT/'ISYS_Diagnostic_Agent_Patch_History.docx'), 'markdown':str(OUT/'ISYS_Diagnostic_Agent_Patch_History.md'), 'planned_pages':len(pages)}))
